"""
data_loader.py
==============
Import robuste du classeur Excel. Lit chaque onglet attendu en DataFrame brut
et signale les onglets manquants. Ne fait AUCUNE interprétation métier : c'est
le rôle de data_cleaning.py.
"""

from __future__ import annotations

import pandas as pd
from openpyxl import load_workbook

from .config import REQUIRED_SHEETS, SHEET_DUREE, SHEET_DIST


class ImportError_(Exception):
    """Erreur bloquante d'import."""


def lire_classeur(chemin_ou_buffer) -> dict[str, pd.DataFrame]:
    """
    Lit toutes les feuilles d'un classeur Excel.
    Retourne {nom_feuille: DataFrame}. Les feuilles absentes ne sont pas créées.
    """
    try:
        wb = load_workbook(chemin_ou_buffer, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ImportError_(f"Impossible d'ouvrir le fichier Excel : {exc}") from exc

    feuilles: dict[str, pd.DataFrame] = {}
    for nom in wb.sheetnames:
        ws = wb[nom]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            feuilles[nom] = pd.DataFrame()
            continue
        # Pour les matrices, l'en-tête est sur la 1re ligne mais la 1re cellule
        # est vide -> on garde tel quel, le cleaning s'en charge.
        df = pd.DataFrame(rows[1:], columns=[_col(c, i) for i, c in enumerate(rows[0])])
        feuilles[nom] = df
    wb.close()
    return feuilles


def _col(valeur, index: int) -> str:
    if valeur is None:
        return f"_col{index}"
    return str(valeur).strip()


def controler_onglets(feuilles: dict[str, pd.DataFrame]) -> list[str]:
    """Retourne la liste des onglets obligatoires manquants."""
    return [s for s in REQUIRED_SHEETS if s not in feuilles]


def matrices_brutes(feuilles: dict[str, pd.DataFrame]) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Retourne (durée, distance) telles quelles pour contrôle ultérieur."""
    return feuilles.get(SHEET_DUREE, pd.DataFrame()), feuilles.get(SHEET_DIST, pd.DataFrame())
